"""Load a resume into Claude message content blocks."""

from __future__ import annotations

import base64
from pathlib import Path

SUPPORTED = (".pdf", ".docx", ".txt", ".md")


def load_resume(path: str | Path) -> list[dict]:
    """Return content blocks describing the resume.

    PDFs are sent to Claude natively (it reads layout, tables, etc.).
    DOCX/TXT/MD are converted to plain text.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
        return [
            {
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf", "data": data},
                "title": path.name,
            }
        ]
    if suffix == ".docx":
        text = _docx_text(path)
    elif suffix in (".txt", ".md"):
        text = path.read_text(encoding="utf-8")
    else:
        raise ValueError(f"Unsupported resume format {suffix!r}; use one of {', '.join(SUPPORTED)}")

    if not text.strip():
        raise ValueError(f"Resume {path} appears to be empty")
    return [{"type": "text", "text": f"<resume>\n{text}\n</resume>"}]


def _docx_text(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
